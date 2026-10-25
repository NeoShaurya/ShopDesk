"""Excel (.xlsx) export for any report or on-screen table, for analysis or to hand over to the CA."""
import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QFileDialog, QMessageBox, QPushButton

_NUM = re.compile(r"^-?\d[\d,]*(\.\d+)?%?$")
HEAD = PatternFill("solid", fgColor="CBD5E1")
FOOT = PatternFill("solid", fgColor="E2E8F0")


def _cell(value, numeric):
    """(value, number format). Only right-aligned columns are turned into real numbers, so text such as
    HSN codes with leading zeros, GSTINs and invoice numbers is never altered."""
    if not numeric:
        return value, None
    if isinstance(value, bool):
        return value, None
    if isinstance(value, int):
        return value, "#,##0"
    if isinstance(value, float):
        return value, "#,##0.00"
    s = str(value).strip()
    if _NUM.match(s):
        pct = s.endswith("%")
        body = s.rstrip("%").replace(",", "")
        n = float(body)
        if pct:
            return n / 100, "0%"
        return (int(n), "#,##0") if "." not in body else (n, "#,##0.00")
    return value, None


def _sheet_name(title, used):
    name = re.sub(r"[\[\]:*?/\\]", "", title)[:31] or "Sheet"
    base, n = name, 2
    while name.lower() in used:
        name = f"{base[:28]} {n}"
        n += 1
    used.add(name.lower())
    return name


def _write_sheet(ws, rep):
    aligns = rep["aligns"]
    ws.append([rep["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    if rep.get("subtitle"):
        ws.append([rep["subtitle"]])
        ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="64748B")
    ws.append([])
    head_row = ws.max_row + 1
    ws.append(list(rep["headers"]))
    for c in range(1, len(rep["headers"]) + 1):
        cell = ws.cell(row=head_row, column=c)
        cell.font, cell.fill = Font(bold=True), HEAD
        cell.alignment = Alignment(horizontal="right" if aligns[c - 1] == "r" else "left")

    def put(cells, bold=False, fill=None):
        ws.append([None] * len(cells))
        r = ws.max_row
        for c, v in enumerate(cells, 1):
            val, fmt = _cell(v, aligns[c - 1] == "r")
            cell = ws.cell(row=r, column=c, value=val if val != "" else None)
            if fmt:
                cell.number_format = fmt
            if aligns[c - 1] == "r":
                cell.alignment = Alignment(horizontal="right")
            if bold:
                cell.font = Font(bold=True)
            if fill:
                cell.fill = fill

    for row in rep["rows"]:
        put(row)
    last_data = ws.max_row
    for row in rep.get("footer", []):
        put(row, bold=True, fill=FOOT)
    ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    if rep["rows"]:
        ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(rep['headers']))}{last_data}"
    for c in range(1, len(rep["headers"]) + 1):
        longest = max([len(str(rep["headers"][c - 1]))] + [len(str(r[c - 1])) for r in rep["rows"] + rep.get("footer", [])
                                                           if c - 1 < len(r)])
        ws.column_dimensions[get_column_letter(c)].width = min(max(longest + 2, 10), 48)


def write_xlsx(path, reps):
    """reps: list of report dicts (title, subtitle, headers, aligns, rows, footer). One sheet each."""
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for rep in reps:
        _write_sheet(wb.create_sheet(_sheet_name(rep.get("sheet") or rep["title"], used)), rep)
    wb.save(path)


def table_to_rep(table, title):
    """Turns an on-screen QTableWidget into the same report shape."""
    cols = [c for c in range(table.columnCount())
            if table.horizontalHeaderItem(c) and table.horizontalHeaderItem(c).text().strip()]
    rows = []
    for r in range(table.rowCount()):
        rows.append([(table.item(r, c).text() if table.item(r, c) else "") for c in cols])
    aligns = ""
    for c in cols:
        it = table.item(0, c) if table.rowCount() else None
        aligns += "r" if it is not None and (it.textAlignment() & Qt.AlignRight) else "l"
    return {"title": title, "subtitle": f"Exported {date.today():%d %b %Y}",
            "headers": [table.horizontalHeaderItem(c).text() for c in cols], "aligns": aligns, "rows": rows, "footer": []}


def save_dialog(parent, reps, default_name):
    docs = Path.home() / "Documents"
    start = (docs if docs.exists() else Path.home()) / f"{default_name}_{date.today():%Y-%m-%d}.xlsx"
    path, _ = QFileDialog.getSaveFileName(parent, "Export to Excel", str(start), "Excel files (*.xlsx)")
    if not path:
        return None
    if not path.lower().endswith(".xlsx"):
        path += ".xlsx"
    try:
        write_xlsx(path, reps)
    except PermissionError:
        QMessageBox.warning(parent, "Could not save", "That file is open in Excel, or the folder can't be written to.\n"
                            "Close the file (or choose another place) and try again.")
        return None
    except OSError as e:
        QMessageBox.warning(parent, "Could not save", str(e))
        return None
    QMessageBox.information(parent, "Exported", f"Saved to:\n{path}")
    return path


class ExcelButton(QPushButton):
    """'Export Excel' button. `get_reps()` is called on every click so it always exports what is on screen."""

    def __init__(self, get_reps, name, parent=None):
        super().__init__("Export Excel", parent)
        self.setStyleSheet("QPushButton{background:white;color:#0f172a;border:1px solid #cbd5e1;border-radius:6px;"
                           "padding:8px 16px;}QPushButton:hover{background:#f1f5f9;}")
        self.clicked.connect(lambda: save_dialog(self, get_reps(), name))
